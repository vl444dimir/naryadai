from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base, get_db
from app.models import Assessment, Brigade, Equipment, EquipmentDowntime, FaultCode, Material, MaterialNorm, MaterialUsage, OrderEvent as OrderEventModel, Section, Shift, ShiftMember, Submission, User, WorkOrder

from app.services.reports import ReportPeriod, as_utc, export_workbook, in_period, interval_seconds, merged_intervals, order_report, period_report, rating_components, safe_excel_cell


NOW = datetime(2026, 10, 8, 12, tzinfo=timezone.utc)


def OrderEvent(**values):
    # Immutable fixture identities at the event time, just as the API writer records.
    details = dict(values.get("details") or {})
    actor = values.get("actor_id")
    if actor:
        details.setdefault("actor_role", "master" if actor == "master" else "worker")
        details.setdefault("actor_brigade_id", {"a": "ba", "b": "bb"}.get(actor))
    values["details"] = details
    return OrderEventModel(**values)


def test_period_is_half_open_and_normalizes_timezone():
    start, end = NOW - timedelta(hours=1), NOW
    assert in_period(start, start, end)
    assert not in_period(end, start, end)
    assert as_utc(NOW.replace(tzinfo=None)) == NOW


def test_downtime_clips_unions_and_caps_open_intervals_at_now():
    start, end = NOW - timedelta(hours=3), NOW + timedelta(hours=2)
    intervals = [(start - timedelta(hours=1), start + timedelta(hours=1)), (start + timedelta(minutes=30), start + timedelta(hours=2)), (NOW - timedelta(minutes=30), None)]
    merged = merged_intervals(intervals, start, end, NOW)
    assert interval_seconds(merged) == 2.5 * 3600


def test_rating_does_not_invent_missing_seven_day_evidence():
    report = rating_components({"quality": 1, "on_time": 1, "first_pass": None, "complexity": 1, "acceptance": 1})
    assert report["earned_points"] == 85
    assert report["available_weight"] == 85
    assert report["provisional"] is True
    assert report["components"]["first_pass"]["points"] is None
    assert rating_components({})["score"] is None


def test_xlsx_export_keeps_formula_like_user_text_literal():
    from openpyxl import load_workbook

    untrusted = ["=HYPERLINK(\"https://example.org\")", " +SUM(A1:A2)", "\t@SUM(A1)", "\r-3+7"]
    content = export_workbook({"Наряды": [{"Комментарий": text, "Количество": -2} for text in untrusted]})
    workbook = load_workbook(BytesIO(content))
    for row, text in enumerate(untrusted, 2):
        assert workbook.active.cell(row, 1).data_type == "s"
        assert workbook.active.cell(row, 1).value == "'" + text.replace("\r", "\n")
        assert workbook.active.cell(row, 2).value == -2
    assert safe_excel_cell("Обычный комментарий") == "Обычный комментарий"


@pytest.fixture
def report_db():
    engine = create_engine("sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        session.add_all([Section(id="sa", name="Дробление"), Section(id="sb", name="РМЦ"), Brigade(id="ba", name="Бригада А"), Brigade(id="bb", name="Бригада Б")])
        session.flush()
        session.add_all([User(id="master", login="master", full_name="Мастер", password_hash="unused", role="master"), User(id="a", login="a", full_name="Рабочий А", password_hash="unused", role="worker", brigade_id="ba"), User(id="b", login="b", full_name="Рабочий Б", password_hash="unused", role="worker", brigade_id="bb"), Equipment(id="ea", name="К-3", inventory_number="A", section_id="sa", equipment_type="conveyor"), Equipment(id="eb", name="Н-7", inventory_number="B", section_id="sb", equipment_type="pump"), FaultCode(id="fault", code="М-02", name="Подшипник"), Material(id="part", name="Подшипник", unit="шт")])
        session.flush()
        session.add(MaterialNorm(fault_code_id="fault", equipment_type="conveyor", material_id="part", quantity=2))
        session.commit()
        yield session
    engine.dispose()


def make_order(session, number, *, created=None, closed=None, employee="a", section="sa", equipment="ea", brigade="ba", **kwargs):
    captured_brigade = kwargs.pop("executor_brigade_id", session.get(User, employee).brigade_id if employee else None)
    order = WorkOrder(id=f"o{number}", number=number, work_type="unplanned", description="Неисправность", section_id=section, equipment_id=equipment, executor_id=employee, executor_brigade_id=captured_brigade, brigade_id=brigade if employee is None else None, master_id="master", due_at=(created or NOW - timedelta(hours=2)) + timedelta(hours=2), norm_minutes=60, created_at=created or NOW - timedelta(hours=2), status="closed" if closed else "issued", closed_at=closed, synthetic=True, fault_code_id="fault", **kwargs)
    session.add(order)
    session.flush()
    return order


def test_worker_grade_is_own_and_master_has_evidence(report_db):
    order = make_order(report_db, 1, closed=NOW - timedelta(hours=1), master_score=5, result_version=1)
    report_db.add(Assessment(order_id=order.id, input_version=1, verdict="accepted_with_comments", score=2, explanation="Предварительно", created_at=NOW - timedelta(hours=1), good=["Работы описаны"], improve=["Уточнить расход"], engine="rules"))
    report_db.flush()
    worker = order_report(report_db, order.id, report_db.get(User, "a"), NOW)
    master = order_report(report_db, order.id, report_db.get(User, "master"), NOW)
    assert worker["quality"]["score"] == 5
    assert worker["quality"]["source"] == "master"
    assert worker["audience"] == "worker"
    assert "materials" not in worker and "history" not in worker
    assert "history" in master and "photos" in master and master["assessment"]["score"] == 2
    with pytest.raises(LookupError):
        order_report(report_db, order.id, report_db.get(User, "b"), NOW)


def test_filter_scope_is_shared_for_counts_ratings_and_export(report_db):
    make_order(report_db, 1, closed=NOW - timedelta(hours=1), master_score=5)
    make_order(report_db, 2, employee="b", section="sb", equipment="eb", brigade="bb")
    period = ReportPeriod(NOW - timedelta(days=1), NOW, section_id="sa", equipment_id="ea", employee_id="a", brigade_id="ba")
    report = period_report(report_db, report_db.get(User, "master"), period, NOW)
    assert report["order_ids"] == ["o1"]
    assert report["counts"]["issued"] == 1
    assert report["counts"]["completed"] == 1
    assert [row["employee_id"] for row in report["ratings"]["employees"]] == ["a"]
    assert report["ratings"]["employees"][0]["counts"]["pending_7day"] == 1
    with pytest.raises(PermissionError):
        period_report(report_db, report_db.get(User, "a"), ReportPeriod(period.start, period.end, employee_id="b"), NOW)


def test_work_time_excludes_pause_and_downtime_unions_orders(report_db):
    start = NOW - timedelta(hours=2)
    order = make_order(report_db, 1, created=start, closed=NOW - timedelta(minutes=15), submitted_at=NOW - timedelta(minutes=30), working_executor_id="a")
    other = make_order(report_db, 2, created=start)
    for minutes, status in [(0, "in_progress"), (30, "paused"), (90, "in_progress"), (105, "performed"), (110, "closed")]:
        report_db.add(OrderEvent(order_id=order.id, actor_id="a", action=status, to_status=status, occurred_at=start + timedelta(minutes=minutes)))
    report_db.add_all([EquipmentDowntime(equipment_id="ea", order_id=order.id, started_at=start, ended_at=start + timedelta(minutes=80), cause="Подшипник", actor_id="master"), EquipmentDowntime(equipment_id="ea", order_id=other.id, started_at=start + timedelta(minutes=40), ended_at=start + timedelta(minutes=100), cause="Повтор", actor_id="master")])
    report_db.flush()
    report = period_report(report_db, report_db.get(User, "master"), ReportPeriod(start, NOW), NOW)
    assert report["workload"][0]["active_minutes"] == 45
    assert report["workload"][0]["paused_minutes"] == 60
    assert report["downtime"][0]["minutes"] == 100
    assert report["counts"]["equipment_downtime_minutes"] == 100


def test_repeat_is_mature_and_only_master_confirmation_affects_rating(report_db):
    first = make_order(report_db, 1, created=NOW - timedelta(days=14), closed=NOW - timedelta(days=13), master_score=5)
    later = make_order(report_db, 2, created=NOW - timedelta(days=12), closed=NOW - timedelta(days=11), employee="b", brigade="bb", master_score=5)
    report_db.add(OrderEvent(order_id=later.id, actor_id="b", action="comment", occurred_at=NOW - timedelta(days=10), details={"confirmed_repeat_of": first.id, "confirmation_role": "master"}))
    report_db.flush()
    period = ReportPeriod(NOW - timedelta(days=30), NOW, employee_id="a")
    report = period_report(report_db, report_db.get(User, "master"), period, NOW)
    rating = report["ratings"]["employees"][0]
    assert rating["counts"]["mature_7day"] == 1
    assert rating["components"]["first_pass"]["points"] == 15
    report_db.add(OrderEvent(order_id=later.id, actor_id="master", action="confirm_repeat", occurred_at=NOW - timedelta(days=9), details={"confirmed_repeat_of": first.id}))
    report_db.flush()
    updated = period_report(report_db, report_db.get(User, "master"), period, NOW)["ratings"]["employees"][0]
    assert updated["components"]["first_pass"]["points"] == 0
    assert updated["counts"]["confirmed_repeat"] == 1


def test_latest_material_snapshot_and_anomaly_are_exact(report_db):
    order = make_order(report_db, 1, result_version=2)
    report_db.add_all([Submission(order_id=order.id, result_version=1, actor_id="a", work_description="Первый", fault_code_id="fault", submitted_at=NOW - timedelta(hours=1)), Submission(order_id=order.id, result_version=2, actor_id="a", work_description="Исправленный", fault_code_id="fault", submitted_at=NOW - timedelta(minutes=30)), MaterialUsage(order_id=order.id, result_version=1, material_id="part", quantity=100), MaterialUsage(order_id=order.id, result_version=2, material_id="part", quantity=8)])
    report_db.flush()
    report = period_report(report_db, report_db.get(User, "master"), ReportPeriod(NOW - timedelta(days=1), NOW), NOW)
    assert report["materials"]["totals"][0]["quantity"] == 8
    assert report["materials"]["rows"][0]["deviation"] == 6
    signal = next(item for item in report["anomalies"] if item["kind"] == "material_over_norm")
    assert signal["evidence_order_ids"] == [order.id]
    assert signal["association_is_not_blame"] is True
    past = period_report(report_db, report_db.get(User, "master"), ReportPeriod(NOW - timedelta(days=1), NOW - timedelta(minutes=45)), NOW)
    assert past["materials"]["totals"][0]["quantity"] == 100
    assert past["materials"]["rows"][0]["result_version"] == 1


def test_on_time_uses_submission_not_master_waiting(report_db):
    order = make_order(report_db, 1, created=NOW - timedelta(hours=4), closed=NOW - timedelta(minutes=5), submitted_at=NOW - timedelta(hours=3), master_score=5)
    report = period_report(report_db, report_db.get(User, "master"), ReportPeriod(NOW - timedelta(days=1), NOW), NOW)
    assert report["ratings"]["employees"][0]["components"]["on_time"]["points"] == 25
    assert order_report(report_db, order.id, report_db.get(User, "a"), NOW)["timing"]["on_time"] is True


def test_api_auth_period_validation_and_worker_export(report_db):
    from app.api.reports import router
    from app.security import get_current_user
    from openpyxl import load_workbook

    order = make_order(report_db, 1, master_score=4)
    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.dependency_overrides[get_db] = lambda: report_db
    with TestClient(app) as client:
        assert client.get(f"/api/reports/orders/{order.id}").status_code == 401
        app.dependency_overrides[get_current_user] = lambda: report_db.get(User, "a")
        response = client.get(f"/api/reports/orders/{order.id}/export.xlsx")
        assert response.status_code == 200
        assert load_workbook(BytesIO(response.content)).sheetnames == ["Моя оценка"]
        assert client.get("/api/reports/ratings", params={"from_at": "2026-10-07T00:00:00", "to_at": "2026-10-08T00:00:00"}).status_code == 422
        assert client.get("/api/reports/ratings", params={"employee_id": "b"}).status_code == 403


def test_arbitrary_period_longer_than_year_is_allowed(report_db):
    period = ReportPeriod(NOW - timedelta(days=800), NOW)
    report = period_report(report_db, report_db.get(User, "master"), period, NOW)
    assert report["counts"]["issued"] == 0
    assert report["ratings"]["employees"] == []


def test_historical_grade_uses_event_score_not_future_master_field(report_db):
    order = make_order(report_db, 1, created=NOW - timedelta(days=15), closed=NOW - timedelta(days=14), master_score=5)
    report_db.add_all([OrderEvent(order_id=order.id, actor_id="master", action="closed", to_status="closed", occurred_at=NOW - timedelta(days=14), details={"score": 2}), OrderEvent(order_id=order.id, actor_id="master", action="master_review", from_status="closed", to_status="closed", occurred_at=NOW - timedelta(days=2), details={"score": 5})])
    report_db.flush()
    report = period_report(report_db, report_db.get(User, "master"), ReportPeriod(NOW - timedelta(days=20), NOW - timedelta(days=5)), NOW)
    assert report["ratings"]["employees"][0]["quality_average"] == 2
    assert report["ratings"]["employees"][0]["components"]["quality"]["points"] == 16


def test_future_reassignment_deadline_and_rejection_review_do_not_rewrite_past(report_db):
    order = make_order(report_db, 1, created=NOW - timedelta(hours=4), employee="b", rejection_justified=False)
    order.due_at = NOW + timedelta(hours=1)
    rejected = OrderEvent(order_id=order.id, actor_id="a", action="rejected", from_status="issued", to_status="rejected", occurred_at=NOW - timedelta(hours=3), reason="Нет допуска")
    report_db.add(rejected)
    report_db.flush()
    report_db.add_all([OrderEvent(order_id=order.id, actor_id="master", action="reassigned", from_status="rejected", to_status="issued", occurred_at=NOW - timedelta(minutes=30), details={"assignment": {"old_executor_id": "a", "old_brigade_id": None, "executor_id": "b", "brigade_id": None}, "due_at": {"old": (NOW - timedelta(hours=2)).isoformat(), "new": order.due_at.isoformat()}}), OrderEvent(order_id=order.id, actor_id="master", action="rejection_reviewed", from_status="issued", to_status="issued", occurred_at=NOW - timedelta(minutes=10), details={"rejection_event_id": rejected.id, "justified": False})])
    report_db.flush()
    past = period_report(report_db, report_db.get(User, "a"), ReportPeriod(NOW - timedelta(hours=5), NOW - timedelta(hours=1)), NOW)
    assert past["counts"]["overdue"] == 1
    assert past["counts"]["rejected"] == 1
    assert past["ratings"]["employees"][0]["employee_id"] == "a"
    assert past["ratings"]["employees"][0]["counts"]["unreviewed_rejections"] == 1
    current_a = period_report(report_db, report_db.get(User, "a"), ReportPeriod(NOW - timedelta(hours=5), NOW), NOW)
    assert current_a["ratings"]["employees"][0]["counts"]["unjustified_rejections"] == 1
    assert current_a["counts"]["completed"] == 0
    current_b = period_report(report_db, report_db.get(User, "b"), ReportPeriod(NOW - timedelta(hours=5), NOW), NOW)
    assert current_b["counts"]["rejected"] == 0
    assert current_b["ratings"]["employees"][0]["counts"]["unjustified_rejections"] == 0


def test_summary_routes_send_only_computed_numbers_and_enforce_own_rating(report_db, monkeypatch):
    import app.api.reports as api
    from app.security import get_current_user

    make_order(report_db, 1, closed=NOW - timedelta(hours=1), master_score=5)
    payloads = []
    def summarize(facts):
        payloads.append(facts)
        return {"text": "Тестовая сводка", "engine": "rules", "model_status": "disabled"}
    monkeypatch.setattr(api, "generate_summary", summarize)
    app = FastAPI()
    app.include_router(api.router, prefix="/api")
    app.dependency_overrides[get_db] = lambda: report_db
    app.dependency_overrides[get_current_user] = lambda: report_db.get(User, "a")
    params = {"from_at": (NOW - timedelta(days=1)).isoformat(), "to_at": NOW.isoformat()}
    with TestClient(app) as client:
        assert client.post("/api/reports/summary", params=params).status_code == 200
        assert client.post("/api/reports/ratings/a/explanation", params=params).status_code == 200
        assert client.post("/api/reports/ratings/b/explanation", params=params).status_code == 403
    assert len(payloads) == 2
    assert not any(key in payload for payload in payloads for key in ("employee_id", "employee_name", "description", "order_ids", "evidence_order_ids"))
    assert payloads[0]["counts"]["completed"] == 1


def test_reassignment_does_not_give_previous_workers_time_to_new_worker(report_db):
    order = make_order(report_db, 1, employee="b", created=NOW - timedelta(hours=2), closed=NOW - timedelta(minutes=5), working_executor_id="b", submitted_at=NOW - timedelta(minutes=15), master_score=5)
    report_db.add_all([OrderEvent(order_id=order.id, actor_id="a", action="start", from_status="accepted", to_status="in_progress", occurred_at=NOW - timedelta(hours=2)), OrderEvent(order_id=order.id, actor_id="master", action="reassigned", from_status="in_progress", to_status="issued", occurred_at=NOW - timedelta(hours=1), details={"assignment": {"old_executor_id": "a", "old_brigade_id": None, "executor_id": "b", "brigade_id": None}}), OrderEvent(order_id=order.id, actor_id="b", action="start", from_status="accepted", to_status="in_progress", occurred_at=NOW - timedelta(minutes=45)), OrderEvent(order_id=order.id, actor_id="b", action="performed", from_status="in_progress", to_status="performed", occurred_at=NOW - timedelta(minutes=15))])
    report_db.flush()
    period = ReportPeriod(NOW - timedelta(days=1), NOW)
    report = period_report(report_db, report_db.get(User, "master"), period, NOW)
    workload = {row["employee_id"]: row for row in report["workload"]}
    assert workload["a"]["active_minutes"] == 60
    assert workload["b"]["active_minutes"] == 30
    own = period_report(report_db, report_db.get(User, "a"), period, NOW)
    assert own["workload"][0]["employee_id"] == "a"
    assert own["workload"][0]["active_minutes"] == 60
    assert own["counts"]["completed"] == 0
    assert own["materials"]["rows"] == []


def test_shift_uses_actual_period_and_roster_including_carryover(report_db):
    start = NOW - timedelta(hours=8)
    report_db.add_all([Shift(id="oldshift", name="Прошлая смена", starts_at=start - timedelta(hours=8), ends_at=start), Shift(id="newshift", name="Текущая смена", starts_at=start, ends_at=NOW)])
    report_db.flush()
    report_db.add(ShiftMember(shift_id="newshift", user_id="b"))
    order = make_order(report_db, 1, created=start - timedelta(hours=1), closed=NOW - timedelta(hours=1), started_at=start - timedelta(minutes=30), submitted_at=NOW - timedelta(hours=1, minutes=5), master_score=5, shift_id="oldshift")
    report_db.add(EquipmentDowntime(equipment_id="ea", order_id=order.id, started_at=start, ended_at=NOW - timedelta(hours=1), cause="Переходящий ремонт", actor_id="master"))
    report_db.flush()
    report = period_report(report_db, report_db.get(User, "master"), ReportPeriod(start, NOW, shift_id="newshift"), NOW)
    assert report["counts"]["completed"] == 1
    assert report["counts"]["equipment_downtime_minutes"] == 420
    workload = {row["employee_id"]: row for row in report["workload"]}
    assert workload["a"]["active_minutes"] == 415
    assert workload["a"]["in_selected_shift_roster"] is False
    assert workload["b"]["active_minutes"] == 0
    assert workload["b"]["in_selected_shift_roster"] is True
    assert report["roster"]["member_ids"] == ["b"]


def test_historical_brigade_role_materials_and_workload_survive_staff_changes(report_db):
    created = NOW - timedelta(days=15, hours=2)
    submitted = created + timedelta(minutes=55)
    closed = created + timedelta(hours=1)
    order = make_order(report_db, 1, created=created, closed=closed, submitted_at=submitted, working_executor_id="a", result_version=1, master_score=5)
    rejected = OrderEvent(order_id=order.id, actor_id="a", action="rejected", from_status="issued", to_status="rejected", occurred_at=created + timedelta(minutes=1))
    report_db.add(rejected)
    report_db.flush()
    report_db.add_all([Submission(order_id=order.id, result_version=1, actor_id="a", actor_brigade_id="ba", work_description="Замена подшипника", fault_code_id="fault", submitted_at=submitted), MaterialUsage(order_id=order.id, result_version=1, material_id="part", quantity=2), OrderEvent(order_id=order.id, actor_id="master", action="rejection_reviewed", from_status="rejected", to_status="rejected", occurred_at=created + timedelta(minutes=2), details={"rejection_event_id": rejected.id, "justified": False}), OrderEvent(order_id=order.id, actor_id="a", action="started", from_status="accepted", to_status="in_progress", occurred_at=created + timedelta(minutes=5)), OrderEvent(order_id=order.id, actor_id="a", action="performed", from_status="in_progress", to_status="performed", occurred_at=submitted), OrderEvent(order_id=order.id, actor_id="master", action="closed", from_status="ai_review", to_status="closed", occurred_at=closed, details={"score": 5}), User(id="auditor", login="auditor", full_name="Аудитор", password_hash="unused", role="admin")])
    report_db.flush()
    period = ReportPeriod(NOW - timedelta(days=20), NOW - timedelta(days=8), brigade_id="ba")
    before = period_report(report_db, report_db.get(User, "auditor"), period, NOW)
    report_db.get(User, "a").brigade_id = "bb"
    report_db.get(User, "a").role = "master"
    report_db.get(User, "master").role = "worker"
    report_db.flush()
    after = period_report(report_db, report_db.get(User, "auditor"), period, NOW)
    assert before["counts"] == after["counts"]
    assert after["counts"]["completed"] == 1
    assert after["ratings"] == before["ratings"]
    assert after["ratings"]["employees"][0]["counts"]["unjustified_rejections"] == 1
    assert after["ratings"]["employees"][0]["quality_average"] == 5
    assert after["workload"][0]["active_minutes"] == 50
    assert after["materials"]["rows"][0]["brigade_id"] == "ba"
    other_brigade = period_report(report_db, report_db.get(User, "auditor"), ReportPeriod(period.start, period.end, brigade_id="bb"), NOW)
    assert other_brigade["counts"]["completed"] == 0
    assert other_brigade["materials"]["rows"] == []


def test_null_legacy_brigade_never_uses_current_membership(report_db):
    order = make_order(report_db, 1, closed=NOW - timedelta(hours=1), result_version=1, master_score=5, executor_brigade_id=None)
    report_db.add_all([Submission(order_id=order.id, result_version=1, actor_id="a", actor_brigade_id=None, work_description="Старые данные", fault_code_id="fault", submitted_at=NOW - timedelta(hours=1)), MaterialUsage(order_id=order.id, result_version=1, material_id="part", quantity=2)])
    report_db.flush()
    period = ReportPeriod(NOW - timedelta(days=1), NOW)
    unfiltered = period_report(report_db, report_db.get(User, "master"), period, NOW)
    assert unfiltered["materials"]["rows"][0]["brigade_id"] is None
    assert unfiltered["brigade_attribution"]["unknown_execution_orders"] == 1
    assert unfiltered["brigade_attribution"]["estimates_used"] is False
    assert unfiltered["ratings"]["brigades"][0]["brigade_id"] is None
    assert unfiltered["ratings"]["brigades"][0]["counts"]["closed"] == 1
    filtered = period_report(report_db, report_db.get(User, "master"), ReportPeriod(period.start, period.end, brigade_id="ba"), NOW)
    assert filtered["counts"]["completed"] == 0
    assert filtered["ratings"]["brigades"] == []


def test_assignment_and_submission_brigades_have_separate_credit(report_db):
    created = NOW - timedelta(hours=2)
    submitted = NOW - timedelta(minutes=30)
    order = make_order(report_db, 1, created=created, closed=NOW - timedelta(minutes=20), submitted_at=submitted, working_executor_id="a", master_score=5, result_version=1)
    report_db.add_all([OrderEvent(order_id=order.id, actor_id="master", action="issued", to_status="issued", occurred_at=created, details={"executor_id": "a", "executor_brigade_id": "ba", "brigade_id": None}), OrderEvent(order_id=order.id, actor_id="a", action="started", from_status="accepted", to_status="in_progress", occurred_at=created), OrderEvent(order_id=order.id, actor_id="a", action="paused", from_status="in_progress", to_status="paused", occurred_at=created + timedelta(minutes=30), details={"actor_brigade_id": "bb"}), OrderEvent(order_id=order.id, actor_id="a", action="resumed", from_status="paused", to_status="in_progress", occurred_at=created + timedelta(minutes=60), details={"actor_brigade_id": "bb"}), OrderEvent(order_id=order.id, actor_id="a", action="performed", from_status="in_progress", to_status="performed", occurred_at=submitted, details={"actor_brigade_id": "bb"}), Submission(order_id=order.id, result_version=1, actor_id="a", actor_brigade_id="bb", work_description="Ремонт", fault_code_id="fault", submitted_at=submitted), MaterialUsage(order_id=order.id, result_version=1, material_id="part", quantity=2)])
    report_db.flush()
    first = period_report(report_db, report_db.get(User, "master"), ReportPeriod(created, NOW, brigade_id="ba"), NOW)
    second = period_report(report_db, report_db.get(User, "master"), ReportPeriod(created, NOW, brigade_id="bb"), NOW)
    assert first["counts"]["issued"] == 1 and first["counts"]["completed"] == 0
    assert second["counts"]["issued"] == 0 and second["counts"]["completed"] == 1
    assert first["workload"][0]["active_minutes"] == 30
    assert second["workload"][0]["active_minutes"] == 30
    assert first["workload"][0]["closed_orders"] == 0
    assert first["ratings"]["employees"][0]["counts"]["closed"] == 0
    assert first["materials"]["rows"] == []
    assert second["materials"]["rows"][0]["brigade_id"] == "bb"


def test_old_submission_does_not_steal_reassigned_work_or_new_worker_grade(report_db):
    created = NOW - timedelta(hours=2)
    order = make_order(report_db, 1, created=created, employee="b", executor_brigade_id="bb", working_executor_id="b", result_version=1, submitted_at=NOW - timedelta(hours=1))
    order.status = "in_progress"
    report_db.add_all([Submission(order_id=order.id, result_version=1, actor_id="a", actor_brigade_id="ba", work_description="Предыдущий результат", fault_code_id="fault", submitted_at=NOW - timedelta(hours=1)), MaterialUsage(order_id=order.id, result_version=1, material_id="part", quantity=2), Assessment(order_id=order.id, input_version=1, verdict="requires_rework", score=2, explanation="Оценка результата А", engine="llm", created_at=NOW - timedelta(minutes=55)), OrderEvent(order_id=order.id, actor_id="a", action="started", from_status="accepted", to_status="in_progress", occurred_at=created), OrderEvent(order_id=order.id, actor_id="a", action="performed", from_status="in_progress", to_status="performed", occurred_at=NOW - timedelta(hours=1)), OrderEvent(order_id=order.id, actor_id=None, action="review_completed", from_status="performed", to_status="rework", occurred_at=NOW - timedelta(minutes=55)), OrderEvent(order_id=order.id, actor_id="master", action="reassigned", from_status="rework", to_status="issued", occurred_at=NOW - timedelta(minutes=40), details={"assignment": {"old_executor_id": "a", "old_executor_brigade_id": "ba", "old_brigade_id": None, "executor_id": "b", "executor_brigade_id": "bb", "brigade_id": None}}), OrderEvent(order_id=order.id, actor_id="b", action="started", from_status="accepted", to_status="in_progress", occurred_at=NOW - timedelta(minutes=20))])
    report_db.flush()
    second = period_report(report_db, report_db.get(User, "master"), ReportPeriod(created, NOW, employee_id="b", brigade_id="bb"), NOW)
    assert second["counts"]["in_progress"] == 1
    assert second["workload"][0]["employee_id"] == "b"
    assert second["workload"][0]["active_minutes"] == 20
    assert second["materials"]["rows"] == []
    first = period_report(report_db, report_db.get(User, "a"), ReportPeriod(created, NOW), NOW)
    assert first["materials"]["rows"][0]["employee_id"] == "a"
    assert first["materials"]["rows"][0]["brigade_id"] == "ba"
    compact = order_report(report_db, order.id, report_db.get(User, "b"), NOW)
    assert compact["quality"]["score"] is None
    assert compact["quality"]["good"] == []
    assert compact["timing"]["submitted_at"] is None
    assert compact["timing"]["active_minutes"] == 20
