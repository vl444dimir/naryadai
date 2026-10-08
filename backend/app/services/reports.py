"""Deterministic report calculations and exports for work-order evidence."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from typing import Any, Iterable, Mapping

from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.models import Assessment, Brigade, Equipment, EquipmentDowntime, FaultCode, Material, MaterialNorm, MaterialUsage, OrderEvent, Photo, Section, ShiftMember, Submission, User, WorkOrder


RATING_WEIGHTS = {"quality": 40, "on_time": 25, "first_pass": 15, "complexity": 15, "acceptance": 5}
TERMINAL_STATUSES = {"closed", "cancelled"}
UNFILTERED_BRIGADE = object()


@dataclass(frozen=True)
class ReportPeriod:
    start: datetime
    end: datetime
    section_id: str | None = None
    equipment_id: str | None = None
    employee_id: str | None = None
    brigade_id: str | None = None
    shift_id: str | None = None

    def __post_init__(self):
        if self.start.tzinfo is None or self.end.tzinfo is None:
            raise ValueError("Период должен содержать часовой пояс")
        if self.end <= self.start:
            raise ValueError("Конец периода должен быть позже начала")


def _iso(value: datetime | None) -> str | None:
    return as_utc(value).isoformat() if value else None


def _completion_at(order: WorkOrder) -> datetime | None:
    return order.submitted_at or order.closed_at


def _actor_role(event: OrderEvent) -> str | None:
    return (event.details or {}).get("actor_role")


def _actor_brigade(event: OrderEvent) -> str | None:
    return (event.details or {}).get("actor_brigade_id")


def _assignment_brigade(order: Any) -> str | None:
    return order.brigade_id or order.executor_brigade_id


def _execution_brigade(order: Any) -> str | None:
    if getattr(order, "has_submission_snapshot", False):
        return order.submission_actor_brigade_id
    if getattr(order, "has_working_brigade_snapshot", False):
        return order.working_executor_brigade_id
    return _assignment_brigade(order)


def _issued_assignment(order: Any, events: list[OrderEvent]) -> tuple[str | None, str | None]:
    issued = next((event for event in events if event.action in {"issue", "issued"}), None)
    details = issued.details or {} if issued else {}
    if "executor_id" in details or "executor_brigade_id" in details or "brigade_id" in details:
        return details.get("executor_id"), details.get("brigade_id") or details.get("executor_brigade_id")
    reassigned = next((event for event in events if event.action == "reassigned" and isinstance((event.details or {}).get("assignment"), dict)), None)
    if reassigned:
        previous = reassigned.details["assignment"]
        return previous.get("old_executor_id"), previous.get("old_brigade_id") or previous.get("old_executor_brigade_id")
    return order.executor_id, _assignment_brigade(order)


def _rejected_events(events: list[OrderEvent], start: datetime, end: datetime, as_of: datetime) -> list[OrderEvent]:
    return [event for event in events if event.action in {"reject", "rejected"} and in_period(event.occurred_at, start, end) and as_utc(event.occurred_at) <= as_utc(as_of)]


def _score_at(order: WorkOrder, assessment: Assessment | None, events: list[OrderEvent], as_of: datetime) -> tuple[int | None, str]:
    graded = [event for event in events if event.to_status == "closed" and (event.details or {}).get("score") is not None and _actor_role(event) in {"master", "admin"}]
    past = [event for event in graded if as_utc(event.occurred_at) <= as_utc(as_of)]
    if past:
        return past[-1].details["score"], "master"
    if order.master_score is not None and not graded:
        return order.master_score, "master"
    return (assessment.score, assessment.engine) if assessment else (None, "unavailable")


def _rejection_reviews(events: dict[str, list[OrderEvent]], as_of: datetime) -> dict[int, bool]:
    outcomes = {}
    for values in events.values():
        rejected = {event.id for event in values if event.action in {"reject", "rejected"}}
        for event in values:
            details = event.details or {}
            target = details.get("rejection_event_id", details.get("last_rejection_event_id"))
            if event.action == "rejection_reviewed" and target in rejected and _actor_role(event) in {"master", "admin"} and isinstance(details.get("justified"), bool) and as_utc(event.occurred_at) <= as_utc(as_of):
                outcomes[target] = details["justified"]
    return outcomes


def _historical_assignment(employee_id: str):
    return select(OrderEvent.id).where(OrderEvent.order_id == WorkOrder.id, OrderEvent.action == "reassigned", OrderEvent.details["assignment"]["old_executor_id"].as_string() == employee_id).exists()


def _historical_work(employee_id: str):
    return select(OrderEvent.id).where(OrderEvent.order_id == WorkOrder.id, OrderEvent.actor_id == employee_id, OrderEvent.to_status.in_(["in_progress", "paused"])).exists()


def _historical_submission(employee_id: str):
    return select(Submission.id).where(Submission.order_id == WorkOrder.id, Submission.actor_id == employee_id).exists()


def _snapshot_orders(session: Session, orders: list[WorkOrder], events: dict[str, list[OrderEvent]], as_of: datetime) -> list[Any]:
    """Detached values prevent current fields from contaminating historical periods."""
    ids = [order.id for order in orders]
    submissions = list(session.scalars(select(Submission).where(Submission.order_id.in_(ids)).order_by(Submission.submitted_at))) if ids else []
    submission_groups: dict[str, list[Submission]] = defaultdict(list)
    for item in submissions:
        submission_groups[item.order_id].append(item)
    snapshots = []
    for order in orders:
        snapshot = SimpleNamespace(**{column.name: getattr(order, column.name) for column in WorkOrder.__table__.columns})
        future_events = [event for event in events[order.id] if as_utc(event.occurred_at) > as_utc(as_of)]
        for event in reversed(future_events):
            changes = event.details or {}
            assignment = changes.get("assignment")
            if isinstance(assignment, dict):
                snapshot.executor_id = assignment.get("old_executor_id")
                snapshot.executor_brigade_id = assignment.get("old_executor_brigade_id")
                snapshot.brigade_id = assignment.get("old_brigade_id")
            if isinstance(changes.get("due_at"), dict) and changes["due_at"].get("old"):
                snapshot.due_at = datetime.fromisoformat(changes["due_at"]["old"])
            if isinstance(changes.get("priority"), dict) and changes["priority"].get("old"):
                snapshot.priority = changes["priority"]["old"]
        for field in ("accepted_at", "started_at", "submitted_at", "closed_at", "cancelled_at"):
            value = getattr(snapshot, field)
            if value and as_utc(value) > as_utc(as_of):
                setattr(snapshot, field, None)
        accepted = [event for event in events[order.id] if event.action in {"accept", "accepted", "start", "started", "in_progress", "queued", "paused", "pause", "resumed", "resume", "continued", "continue"} and event.actor_id and _actor_role(event) in {"worker", None} and as_utc(event.occurred_at) <= as_utc(as_of)]
        past_reassignments = [event for event in events[order.id] if event.action == "reassigned" and as_utc(event.occurred_at) <= as_utc(as_of)]
        snapshot.working_executor_brigade_id = None
        snapshot.has_working_brigade_snapshot = False
        snapshot.has_submission_snapshot = False
        snapshot.submission_actor_brigade_id = None
        snapshot.submission_actor_id = None
        if accepted and (not past_reassignments or as_utc(accepted[-1].occurred_at) >= as_utc(past_reassignments[-1].occurred_at)):
            snapshot.working_executor_id = accepted[-1].actor_id
            snapshot.working_executor_brigade_id = _actor_brigade(accepted[-1])
            snapshot.has_working_brigade_snapshot = "actor_brigade_id" in (accepted[-1].details or {})
        elif future_events:
            snapshot.working_executor_id = None
        reworks = [event for event in events[order.id] if event.to_status == "rework" and event.from_status != "rework"]
        if reworks:
            snapshot.rework_count = sum(as_utc(event.occurred_at) <= as_utc(as_of) for event in reworks)
        if submission_groups[order.id]:
            valid = [item for item in submission_groups[order.id] if as_utc(item.submitted_at) <= as_utc(as_of)]
            latest = valid[-1] if valid else None
            snapshot.result_version = latest.result_version if latest else 0
            snapshot.fault_code_id = latest.fault_code_id if latest else None
            snapshot.work_description = latest.work_description if latest else None
            snapshot.materials_none = latest.materials_none if latest else False
            snapshot.submitted_at = latest.submitted_at if latest else None
            if latest:
                snapshot.submission_actor_id = latest.actor_id
                snapshot.submission_actor_brigade_id = latest.actor_brigade_id
                if _state_at(snapshot, events[order.id], as_of) in {"performed", "ai_review", "closed"}:
                    snapshot.working_executor_id = latest.actor_id
                    snapshot.has_submission_snapshot = True
        snapshot.status = _state_at(snapshot, events[order.id], as_of)
        snapshots.append(snapshot)
    return snapshots


def scoped_orders(session: Session, user: User, period: ReportPeriod | None = None) -> list[WorkOrder]:
    query = select(WorkOrder)
    if user.role == "worker":
        if period and period.employee_id and period.employee_id != user.id:
            raise PermissionError("Исполнитель может видеть отчёты только по своим нарядам")
        rejected_own = select(OrderEvent.id).where(OrderEvent.order_id == WorkOrder.id, OrderEvent.actor_id == user.id, OrderEvent.action.in_(["reject", "rejected"])).exists()
        query = query.where(or_(WorkOrder.executor_id == user.id, WorkOrder.working_executor_id == user.id, rejected_own, _historical_assignment(user.id), _historical_work(user.id), _historical_submission(user.id)))
    elif user.role not in {"master", "leader", "admin"}:
        raise PermissionError("Нет доступа к отчётам")
    if period:
        for name in ("section_id", "equipment_id"):
            value = getattr(period, name)
            if value:
                query = query.where(getattr(WorkOrder, name) == value)
        if period.brigade_id:
            historical_brigade = select(OrderEvent.id).where(OrderEvent.order_id == WorkOrder.id, or_(OrderEvent.details["actor_brigade_id"].as_string() == period.brigade_id, OrderEvent.details["executor_brigade_id"].as_string() == period.brigade_id, OrderEvent.details["assignment"]["old_brigade_id"].as_string() == period.brigade_id, OrderEvent.details["assignment"]["old_executor_brigade_id"].as_string() == period.brigade_id, OrderEvent.details["assignment"]["executor_brigade_id"].as_string() == period.brigade_id)).exists()
            submitted_in_brigade = select(Submission.id).where(Submission.order_id == WorkOrder.id, Submission.actor_brigade_id == period.brigade_id).exists()
            query = query.where(or_(WorkOrder.brigade_id == period.brigade_id, WorkOrder.executor_brigade_id == period.brigade_id, historical_brigade, submitted_in_brigade))
        if period.employee_id:
            rejected_by_employee = select(OrderEvent.id).where(OrderEvent.order_id == WorkOrder.id, OrderEvent.actor_id == period.employee_id, OrderEvent.action.in_(["reject", "rejected"])).exists()
            query = query.where(or_(WorkOrder.executor_id == period.employee_id, WorkOrder.working_executor_id == period.employee_id, rejected_by_employee, _historical_assignment(period.employee_id), _historical_work(period.employee_id), _historical_submission(period.employee_id)))
        query = query.where(WorkOrder.created_at < period.end)
    return list(session.scalars(query.order_by(WorkOrder.number)))


def _events_by_order(session: Session, ids: list[str]) -> dict[str, list[OrderEvent]]:
    grouped: dict[str, list[OrderEvent]] = defaultdict(list)
    if ids:
        for event in session.scalars(select(OrderEvent).where(OrderEvent.order_id.in_(ids)).order_by(OrderEvent.occurred_at, OrderEvent.id)):
            grouped[event.order_id].append(event)
    return grouped


def _current_assessments(session: Session, orders: list[WorkOrder], as_of: datetime) -> dict[str, Assessment]:
    result = {}
    versions = {order.id: order.result_version for order in orders}
    if versions:
        submissions = list(session.scalars(select(Submission).where(Submission.order_id.in_(versions)).order_by(Submission.submitted_at)))
        submitted_orders = {submission.order_id for submission in submissions}
        for order_id in submitted_orders:
            versions[order_id] = 0
        for submission in submissions:
            if as_utc(submission.submitted_at) <= as_utc(as_of):
                versions[submission.order_id] = submission.result_version
        for assessment in session.scalars(select(Assessment).where(Assessment.order_id.in_(versions), Assessment.created_at <= as_of).order_by(Assessment.created_at)):
            if assessment.input_version == versions[assessment.order_id]:
                result[assessment.order_id] = assessment
    return result


def _state_at(order: WorkOrder, events: list[OrderEvent], at: datetime) -> str:
    state = "issued"
    for event in events:
        if as_utc(event.occurred_at) <= as_utc(at) and event.to_status:
            state = event.to_status
    if order.closed_at and as_utc(order.closed_at) <= as_utc(at):
        return "closed"
    if order.cancelled_at and as_utc(order.cancelled_at) <= as_utc(at):
        return "cancelled"
    if not events and as_utc(order.created_at) <= as_utc(at):
        if order.started_at and as_utc(order.started_at) <= as_utc(at):
            state = "in_progress"
        elif order.accepted_at and as_utc(order.accepted_at) <= as_utc(at):
            state = "accepted"
    return state


def work_intervals(order: WorkOrder, events: list[OrderEvent], end: datetime) -> dict[str, list[tuple[datetime, datetime | None]]]:
    result: dict[str, list[tuple[datetime, datetime | None]]] = defaultdict(list)
    transitions = [(as_utc(event.occurred_at), event.to_status) for event in events if event.to_status and as_utc(event.occurred_at) <= as_utc(end)]
    if not transitions:
        if order.started_at:
            result["in_progress"].append((order.started_at, order.submitted_at or order.closed_at))
        return result
    state, start = "issued", as_utc(order.created_at)
    for occurred, next_state in transitions:
        if occurred >= start and state in {"in_progress", "paused"}:
            result[state].append((start, occurred))
        state, start = next_state, occurred
    if state in {"in_progress", "paused"}:
        result[state].append((start, end))
    return result


def actor_brigade_work_intervals(order: WorkOrder, events: list[OrderEvent], end: datetime) -> dict[tuple[str, str | None], dict[str, list[tuple[datetime, datetime]]]]:
    result = defaultdict(lambda: defaultdict(list))
    transitions = [event for event in events if event.to_status and as_utc(event.occurred_at) <= as_utc(end)]
    if not transitions:
        owner = order.working_executor_id or order.executor_id
        if owner:
            result[(owner, _execution_brigade(order))].update(work_intervals(order, events, end))
        return result
    owner, state, start = order.executor_id, "issued", as_utc(order.created_at)
    brigade = _assignment_brigade(order)
    for event in transitions:
        occurred = as_utc(event.occurred_at)
        if owner and state in {"in_progress", "paused"} and occurred > start:
            result[(owner, brigade)][state].append((start, occurred))
        if event.action == "reassigned":
            assignment = (event.details or {}).get("assignment", {})
            owner = assignment.get("executor_id")
            brigade = assignment.get("brigade_id") or assignment.get("executor_brigade_id")
        elif event.to_status in {"accepted", "queued", "in_progress", "paused"} and (_actor_role(event) == "worker" or _actor_role(event) is None and event.action in {"accept", "accepted", "start", "started", "in_progress", "pause", "paused", "resume", "resumed", "continue", "continued", "queue", "queued"}):
            owner = event.actor_id
            brigade = _actor_brigade(event)
        state, start = event.to_status, occurred
    if owner and state in {"in_progress", "paused"} and as_utc(end) > start:
        result[(owner, brigade)][state].append((start, end))
    return result


def actor_work_intervals(order: WorkOrder, events: list[OrderEvent], end: datetime) -> dict[str, dict[str, list[tuple[datetime, datetime]]]]:
    result = defaultdict(lambda: defaultdict(list))
    for (actor, _), states in actor_brigade_work_intervals(order, events, end).items():
        for state, values in states.items():
            result[actor][state].extend(values)
    return result


def _event_dict(event: OrderEvent) -> dict[str, Any]:
    return {"id": event.id, "actor_id": event.actor_id, "action": event.action, "occurred_at": _iso(event.occurred_at), "from_status": event.from_status, "to_status": event.to_status, "reason": event.reason, "details": event.details}


def _assessment_dict(assessment: Assessment | None) -> dict[str, Any] | None:
    if assessment is None:
        return None
    return {"id": assessment.id, "input_version": assessment.input_version, "verdict": assessment.verdict, "score": assessment.score, "explanation": assessment.explanation, "good": assessment.good, "improve": assessment.improve, "findings": assessment.findings, "details": assessment.details, "engine": assessment.engine, "model": assessment.model, "needs_master_review": assessment.needs_master_review, "created_at": _iso(assessment.created_at)}


def order_report(session: Session, order_id: str, user: User, now: datetime) -> dict[str, Any]:
    order = session.get(WorkOrder, order_id)
    if order is None:
        raise LookupError("Наряд не найден")
    if user.role == "worker" and user.id not in {order.executor_id, order.working_executor_id}:
        raise LookupError("Наряд не найден")
    if user.role not in {"worker", "master", "leader", "admin"}:
        raise PermissionError("Нет доступа к отчёту")
    events = _events_by_order(session, [order.id])[order.id]
    assessment = _current_assessments(session, [order], now).get(order.id)
    submission = session.scalar(select(Submission).where(Submission.order_id == order.id, Submission.result_version == order.result_version))
    work = work_intervals(order, events, now)
    active_seconds = interval_seconds(merged_intervals(work["in_progress"], order.created_at, now, now)) if as_utc(now) > as_utc(order.created_at) else 0
    score, score_source = _score_at(order, assessment, events, now)
    other_result = user.role == "worker" and submission is not None and submission.actor_id != user.id
    if other_result:
        assessment = None
        score, score_source = None, "unavailable"
    if user.role == "worker" and as_utc(now) > as_utc(order.created_at):
        own_intervals = actor_work_intervals(order, events, now).get(user.id, {}).get("in_progress", [])
        active_seconds = interval_seconds(merged_intervals(own_intervals, order.created_at, now, now))
    common = {
        "order_id": order.id, "number": order.number, "status": order.status,
        "quality": {"score": score, "score_scale": [1, 5], "source": score_source, "verdict": None if other_result else order.master_verdict or (assessment.verdict if assessment else None), "explanation": "Оценка вашего результата ещё не готова" if other_result else order.master_comment or (assessment.explanation if assessment else "Оценка ещё не готова"), "good": assessment.good if assessment else [], "improve": assessment.improve if assessment else [], "needs_master_review": True if other_result else False if order.master_verdict else assessment.needs_master_review if assessment else True},
        "timing": {"norm_minutes": order.norm_minutes, "active_minutes": round(active_seconds / 60, 2), "due_at": _iso(order.due_at), "started_at": _iso(order.started_at), "submitted_at": None if other_result else _iso(order.submitted_at), "closed_at": _iso(order.closed_at), "on_time": as_utc(_completion_at(order)) <= as_utc(order.due_at) if _completion_at(order) and not other_result else None, "on_time_basis": "submitted_at; closed_at fallback"},
        "synthetic": order.synthetic,
    }
    if user.role == "worker":
        return {**common, "audience": "worker"}
    equipment = session.get(Equipment, order.equipment_id)
    section = session.get(Section, order.section_id)
    material_lookup = {item.id: item for item in session.scalars(select(Material))}
    usages = list(session.scalars(select(MaterialUsage).where(MaterialUsage.order_id == order.id, MaterialUsage.result_version == order.result_version)))
    photos = list(session.scalars(select(Photo).where(Photo.order_id == order.id).order_by(Photo.uploaded_at)))
    downtimes = list(session.scalars(select(EquipmentDowntime).where(EquipmentDowntime.order_id == order.id)))
    downtime_start = min((as_utc(item.started_at) for item in downtimes), default=as_utc(order.created_at))
    merged = merged_intervals([(item.started_at, item.ended_at) for item in downtimes], downtime_start, now, now) if as_utc(now) > downtime_start else []
    return {
        **common, "audience": "master", "description": order.description, "comment": order.comment,
        "work_type": order.work_type, "priority": order.priority, "section_id": order.section_id, "section_name": section.name if section else None,
        "equipment_id": order.equipment_id, "equipment_name": equipment.name if equipment else None,
        "executor_id": order.executor_id, "executor_brigade_id": order.executor_brigade_id, "working_executor_id": order.working_executor_id, "brigade_id": order.brigade_id, "result_actor_brigade_id": submission.actor_brigade_id if submission else None, "master_id": order.master_id,
        "created_at": _iso(order.created_at), "work_description": order.work_description, "fault_code_id": order.fault_code_id, "materials_none": order.materials_none,
        "result_version": order.result_version, "rework_count": order.rework_count, "history": [_event_dict(event) for event in events],
        "materials": [{"id": item.id, "material_id": item.material_id, "name": material_lookup[item.material_id].name, "unit": material_lookup[item.material_id].unit, "quantity": float(item.quantity), "result_version": item.result_version} for item in usages],
        "photos": [{"id": photo.id, "kind": photo.kind, "url": f"/api/photos/{photo.id}/file", "author_id": photo.author_id, "uploaded_at": _iso(photo.uploaded_at), "captured_at": _iso(photo.captured_at), "result_version": photo.result_version, "source": photo.source} for photo in photos if photo.kind == "before" or photo.result_version == order.result_version],
        "assessment": _assessment_dict(assessment),
        "downtime": {"minutes": round(interval_seconds(merged) / 60, 2), "intervals": [{"id": item.id, "started_at": _iso(item.started_at), "ended_at": _iso(item.ended_at), "cause": item.cause} for item in downtimes], "basis": "explicit_equipment_downtime"},
    }


def _relevant(order: WorkOrder, events: list[OrderEvent], period: ReportPeriod) -> bool:
    if in_period(order.created_at, period.start, period.end) or in_period(order.closed_at, period.start, period.end) or in_period(order.submitted_at, period.start, period.end):
        return True
    if any(in_period(event.occurred_at, period.start, period.end) for event in events):
        return True
    terminal = order.closed_at or order.cancelled_at
    return as_utc(order.created_at) < as_utc(period.end) and (terminal is None or as_utc(terminal) >= as_utc(period.start))


def _confirmed_repeats(orders: list[WorkOrder], events: dict[str, list[OrderEvent]], now: datetime) -> set[str]:
    indexed = {order.id: order for order in orders}
    result = set()
    for later in orders:
        for event in events[later.id]:
            previous_ids = (event.details or {}).get("confirmed_repeat_of", [])
            if isinstance(previous_ids, str):
                previous_ids = [previous_ids]
            actor = _actor_role(event)
            if not isinstance(previous_ids, list) or as_utc(event.occurred_at) > as_utc(now) or actor not in {"master", "admin"}:
                continue
            for previous_id in previous_ids:
                previous = indexed.get(previous_id)
                if previous and previous.closed_at and previous.equipment_id == later.equipment_id and previous.fault_code_id and previous.fault_code_id == later.fault_code_id and 0 <= (as_utc(later.created_at) - as_utc(previous.closed_at)).total_seconds() <= 7 * 86400:
                    result.add(previous_id)
    return result


def _rating_row(orders: list[WorkOrder], assessments: dict[str, Assessment], events: dict[str, list[OrderEvent]], confirmed_repeats: set[str], period: ReportPeriod, now: datetime, rejection_reviews: dict[int, bool], employee_id: str | None = None, brigade_id: Any = UNFILTERED_BRIGADE) -> dict[str, Any]:
    closed = [order for order in orders if in_period(order.closed_at, period.start, period.end) and as_utc(order.closed_at) <= as_utc(now) and (employee_id is None or (order.working_executor_id or order.executor_id) == employee_id) and (brigade_id is UNFILTERED_BRIGADE or _execution_brigade(order) == brigade_id)]
    scores = [_score_at(order, assessments.get(order.id), events[order.id], now)[0] for order in closed]
    known_scores = [score for score in scores if score is not None]
    mature = [order for order in closed if as_utc(order.closed_at) + timedelta(days=7) <= as_utc(now)]
    first_pass = [order for order in mature if order.rework_count == 0 and order.id not in confirmed_repeats]
    rejection_events = [event for order in orders for event in _rejected_events(events[order.id], period.start, period.end, now) if (employee_id is None or event.actor_id == employee_id) and (brigade_id is UNFILTERED_BRIGADE or _actor_brigade(event) == brigade_id)]
    justified_order_ids = {event.order_id for event in rejection_events if rejection_reviews.get(event.id) is True}
    assigned_minutes = sum(order.norm_minutes for order in orders if _state_at(order, events[order.id], now) != "cancelled" and (order.id not in justified_order_ids or order in closed))
    completed_minutes = sum(order.norm_minutes for order in closed)
    unjustified = sum(rejection_reviews.get(event.id) is False for event in rejection_events)
    rejection_unknown = sum(event.id not in rejection_reviews for event in rejection_events)
    on_time = sum(1 for order in closed if as_utc(_completion_at(order)) <= as_utc(order.due_at))
    quality_average = sum(known_scores) / len(known_scores) if known_scores else None
    component_result = rating_components({"quality": quality_average / 5 if quality_average is not None else None, "on_time": on_time / len(closed) if closed else None, "first_pass": len(first_pass) / len(mature) if mature else None, "complexity": completed_minutes / assigned_minutes if assigned_minutes else None, "acceptance": 1 - unjustified / len(orders) if orders and not rejection_unknown else None})
    return {
        **component_result,
        "counts": {"assigned": len(orders), "closed": len(closed), "scored": len(known_scores), "on_time": on_time, "mature_7day": len(mature), "pending_7day": len(closed) - len(mature), "first_pass_mature": len(first_pass), "rework": sum(order.rework_count > 0 for order in closed), "confirmed_repeat": sum(order.id in confirmed_repeats for order in mature), "unjustified_rejections": unjustified, "unreviewed_rejections": rejection_unknown},
        "quality_average": round(quality_average, 2) if quality_average is not None else None,
        "complexity": {"completed_norm_minutes": completed_minutes, "assigned_norm_minutes": assigned_minutes},
        "evidence_order_ids": [order.id for order in closed],
        "explanation": "Q — средняя оценка мастера либо последней проверки / 5; T — доля закрытых нарядов, сданных в срок (submitted_at, иначе closed_at); F — доля без доработки и подтверждённого повторного отказа среди завершённых не менее 7 дней назад; C — закрытые нормативные минуты / назначенные нормативные минуты; A — доля без признанных мастером необоснованных отказов. Неподтверждённая связь повторов не снижает рейтинг. Неполные компоненты отмечены как предварительные.",
    }


def period_report(session: Session, user: User, period: ReportPeriod, now: datetime) -> dict[str, Any]:
    now = min(as_utc(now), as_utc(period.end))
    scoped = [order for order in scoped_orders(session, user, period) if as_utc(order.created_at) <= now]
    events = _events_by_order(session, [order.id for order in scoped])
    scoped = _snapshot_orders(session, scoped, events, now)
    scoped_ids = [order.id for order in scoped]
    overlapping_downtimes = list(session.scalars(select(EquipmentDowntime).where(EquipmentDowntime.order_id.in_(scoped_ids), EquipmentDowntime.started_at < period.end, or_(EquipmentDowntime.ended_at.is_(None), EquipmentDowntime.ended_at > period.start)))) if scoped_ids else []
    downtime_ids = {record.order_id for record in overlapping_downtimes}
    orders = [order for order in scoped if _relevant(order, events[order.id], period) or order.id in downtime_ids]
    all_users = {item.id: item for item in session.scalars(select(User))}
    employees = all_users
    target_employee = user.id if user.role == "worker" else period.employee_id
    if target_employee:
        orders = [order for order in orders if target_employee in {order.executor_id, order.working_executor_id} or (order.submission_actor_id == target_employee and in_period(order.submitted_at, period.start, period.end)) or any(event.actor_id == target_employee for event in _rejected_events(events[order.id], period.start, period.end, now)) or any(merged_intervals(values, period.start, period.end, now) for values in actor_work_intervals(order, events[order.id], now).get(target_employee, {}).values())]
    operational_orders = [order for order in orders if not target_employee or (order.working_executor_id or order.executor_id) == target_employee]
    if period.brigade_id:
        operational_orders = [order for order in operational_orders if period.brigade_id in {_assignment_brigade(order), _execution_brigade(order)} or any(_actor_brigade(event) == period.brigade_id and in_period(event.occurred_at, period.start, period.end) for event in events[order.id])]
    operational_ids = {order.id for order in operational_orders}
    ids = [order.id for order in orders]
    assessments = _current_assessments(session, orders, now)
    rejection_reviews = _rejection_reviews(events, now)
    brigades = {item.id: item for item in session.scalars(select(Brigade))}
    equipment = {item.id: item for item in session.scalars(select(Equipment))}
    sections = {item.id: item for item in session.scalars(select(Section))}
    material_lookup = {item.id: item for item in session.scalars(select(Material))}
    faults = {item.id: item for item in session.scalars(select(FaultCode))}
    history_equipment_ids = {order.equipment_id for order in scoped}
    history_orders = list(session.scalars(select(WorkOrder).where(WorkOrder.equipment_id.in_(history_equipment_ids), WorkOrder.created_at < period.end, WorkOrder.created_at <= now))) if history_equipment_ids else []
    history_events = _events_by_order(session, [order.id for order in history_orders])
    history_orders = _snapshot_orders(session, history_orders, history_events, now)
    confirmed = _confirmed_repeats(history_orders, history_events, now)
    relevant_by_employee: dict[str, list[WorkOrder]] = defaultdict(list)
    relevant_by_brigade: dict[str, list[WorkOrder]] = defaultdict(list)
    workload_by_employee: dict[str, dict[str, list[tuple[datetime, datetime | None]]]] = defaultdict(lambda: defaultdict(list))
    for order in operational_orders:
        owner = order.working_executor_id or order.executor_id
        if owner:
            relevant_by_employee[owner].append(order)
        for brigade in {_assignment_brigade(order), _execution_brigade(order)}:
            if (not period.brigade_id or brigade == period.brigade_id) and order not in relevant_by_brigade[brigade]:
                relevant_by_brigade[brigade].append(order)
    for order in orders:
        for (actor, actor_brigade), states in actor_brigade_work_intervals(order, events[order.id], now).items():
            if actor in employees and (not target_employee or actor == target_employee) and (not period.brigade_id or actor_brigade == period.brigade_id):
                for state, values in states.items():
                    clipped = merged_intervals(values, period.start, period.end, now)
                    workload_by_employee[actor][state].extend(clipped)
                    if clipped and order not in relevant_by_employee[actor]:
                        relevant_by_employee[actor].append(order)
                    if clipped and order not in relevant_by_brigade[actor_brigade]:
                        relevant_by_brigade[actor_brigade].append(order)
        for event in _rejected_events(events[order.id], period.start, period.end, now):
            if event.actor_id in employees and (not target_employee or event.actor_id == target_employee) and (not period.brigade_id or _actor_brigade(event) == period.brigade_id) and order not in relevant_by_employee[event.actor_id]:
                relevant_by_employee[event.actor_id].append(order)
            actor_brigade = _actor_brigade(event)
            if (not period.brigade_id or actor_brigade == period.brigade_id) and order not in relevant_by_brigade[actor_brigade]:
                relevant_by_brigade[actor_brigade].append(order)
    rating_rows = [{"employee_id": employee_id, "employee_name": employees[employee_id].full_name if employee_id in employees else employee_id, **_rating_row(items, assessments, events, confirmed, period, now, rejection_reviews, employee_id, brigade_id=period.brigade_id if period.brigade_id else UNFILTERED_BRIGADE)} for employee_id, items in relevant_by_employee.items()]
    rating_rows.sort(key=lambda row: (row["score"] is None, -(row["score"] or 0), row["employee_id"]))
    brigade_rows = [{"brigade_id": brigade_id, "brigade_name": brigades[brigade_id].name if brigade_id in brigades else "Историческая бригада неизвестна", **_rating_row(items, assessments, events, confirmed, period, now, rejection_reviews, brigade_id=brigade_id)} for brigade_id, items in relevant_by_brigade.items()]
    workload = []
    roster_ids = set(session.scalars(select(ShiftMember.user_id).where(ShiftMember.shift_id == period.shift_id))) if period.shift_id else set()
    visible_roster_ids = {member for member in roster_ids if (not target_employee or member == target_employee)}
    workload_ids = set(relevant_by_employee) | (visible_roster_ids if not period.brigade_id else set())
    for employee_id in sorted(workload_ids):
        assigned = relevant_by_employee[employee_id]
        active = merged_intervals(workload_by_employee[employee_id]["in_progress"], period.start, period.end, now)
        waiting = merged_intervals(workload_by_employee[employee_id]["paused"], period.start, period.end, now)
        workload.append({"employee_id": employee_id, "employee_name": all_users[employee_id].full_name if employee_id in all_users else employee_id, "assigned_orders": len(assigned), "closed_orders": sum(in_period(order.closed_at, period.start, period.end) and (order.working_executor_id or order.executor_id) == employee_id and (not period.brigade_id or _execution_brigade(order) == period.brigade_id) for order in assigned), "active_minutes": round(interval_seconds(active) / 60, 2), "paused_minutes": round(interval_seconds(waiting) / 60, 2), "basis": "status_event_intervals", "waiting_is_not_blame": True, "in_selected_shift_roster": employee_id in roster_ids if period.shift_id else None})
    order_map = {order.id: order for order in orders}
    downtime_records = [record for record in overlapping_downtimes if record.order_id in operational_ids and (not period.brigade_id or _execution_brigade(order_map[record.order_id]) == period.brigade_id)]
    downtime_groups: dict[str, list[EquipmentDowntime]] = defaultdict(list)
    for record in downtime_records:
        downtime_groups[record.equipment_id].append(record)
    downtime_rows = []
    for equipment_id, records in downtime_groups.items():
        all_intervals = merged_intervals([(record.started_at, record.ended_at) for record in records], period.start, period.end, now)
        planned_intervals = merged_intervals([(record.started_at, record.ended_at) for record in records if order_map[record.order_id].work_type == "planned"], period.start, period.end, now)
        unplanned_intervals = merged_intervals([(record.started_at, record.ended_at) for record in records if order_map[record.order_id].work_type == "unplanned"], period.start, period.end, now)
        total, planned, unplanned = interval_seconds(all_intervals), interval_seconds(planned_intervals), interval_seconds(unplanned_intervals)
        unit = equipment.get(equipment_id)
        downtime_rows.append({"equipment_id": equipment_id, "equipment_name": unit.name if unit else equipment_id, "section_id": unit.section_id if unit else None, "minutes": round(total / 60, 2), "planned_minutes": round(planned / 60, 2), "unplanned_minutes": round(unplanned / 60, 2), "overlap_minutes": round((planned + unplanned - total) / 60, 2), "unplanned_share": round(unplanned / total, 4) if total else None, "intervals": [{"started_at": _iso(left), "ended_at": _iso(right)} for left, right in all_intervals], "causes": list(dict.fromkeys(record.cause for record in records)), "evidence_order_ids": list(dict.fromkeys(record.order_id for record in records)), "basis": "explicit_equipment_downtime_union"})
    downtime_rows.sort(key=lambda row: (-row["minutes"], row["equipment_id"]))
    submissions = list(session.scalars(select(Submission).where(Submission.order_id.in_(ids), Submission.submitted_at <= now).order_by(Submission.submitted_at))) if ids else []
    latest_submission = {submission.order_id: submission for submission in submissions}
    material_rows = []
    norms = {(item.fault_code_id, item.equipment_type, item.material_id): float(item.quantity) for item in session.scalars(select(MaterialNorm))}
    for usage in session.scalars(select(MaterialUsage).where(MaterialUsage.order_id.in_(ids))) if ids else []:
        order = order_map[usage.order_id]
        submitted = latest_submission.get(order.id)
        if submitted is None or usage.result_version != submitted.result_version or not in_period(submitted.submitted_at, period.start, period.end) or (target_employee and submitted.actor_id != target_employee) or (period.brigade_id and submitted.actor_brigade_id != period.brigade_id):
            continue
        material = material_lookup[usage.material_id]
        unit = equipment[order.equipment_id]
        norm = norms.get((submitted.fault_code_id, unit.equipment_type, usage.material_id))
        quantity = float(usage.quantity)
        owner = order.working_executor_id or order.executor_id
        material_rows.append({"order_id": order.id, "number": order.number, "material_id": material.id, "material_name": material.name, "unit": material.unit, "quantity": quantity, "norm_quantity": norm, "deviation": round(quantity - norm, 3) if norm is not None else None, "section_id": order.section_id, "equipment_id": order.equipment_id, "employee_id": submitted.actor_id, "brigade_id": submitted.actor_brigade_id, "brigade_basis": "submission_actor_snapshot", "result_version": usage.result_version, "submitted_at": _iso(submitted.submitted_at)})
    material_totals: dict[str, dict[str, Any]] = {}
    for row in material_rows:
        item = material_totals.setdefault(row["material_id"], {"material_id": row["material_id"], "material_name": row["material_name"], "unit": row["unit"], "quantity": 0.0, "order_ids": []})
        item["quantity"] = round(item["quantity"] + row["quantity"], 3)
        item["order_ids"].append(row["order_id"])
    anomalies = _anomalies(operational_orders, history_orders if user.role != "worker" else operational_orders, downtime_rows, material_rows, equipment, sections, faults, period, now)
    issued = [order for order in orders if in_period(order.created_at, period.start, period.end) and as_utc(order.created_at) <= now and (not target_employee or _issued_assignment(order, events[order.id])[0] == target_employee) and (not period.brigade_id or _issued_assignment(order, events[order.id])[1] == period.brigade_id)]
    closed = [order for order in operational_orders if in_period(order.closed_at, period.start, period.end) and as_utc(order.closed_at) <= now and (not period.brigade_id or _execution_brigade(order) == period.brigade_id)]
    rejected_ids = {order.id for order in orders if any((not target_employee or event.actor_id == target_employee) and (not period.brigade_id or _actor_brigade(event) == period.brigade_id) for event in _rejected_events(events[order.id], period.start, period.end, now))}
    current_orders = [order for order in operational_orders if not period.brigade_id or _execution_brigade(order) == period.brigade_id]
    overdue = [order for order in current_orders if as_utc(order.due_at) < now and _state_at(order, events[order.id], now) not in TERMINAL_STATUSES | {"performed", "ai_review"}]
    counts = {"issued": len(issued), "completed": len(closed), "performed": sum(in_period(order.submitted_at, period.start, period.end) and as_utc(order.submitted_at) <= now for order in current_orders), "overdue": len(overdue), "rejected": len(rejected_ids), "in_progress": sum(_state_at(order, events[order.id], now) == "in_progress" for order in current_orders), "pending_approval": sum(_state_at(order, events[order.id], now) in {"performed", "ai_review"} for order in current_orders), "closed_late": sum(as_utc(order.closed_at) > as_utc(order.due_at) for order in closed), "equipment_in_downtime": sum(any(as_utc(record.started_at) <= now and (record.ended_at is None or as_utc(record.ended_at) > now) for record in records) for records in downtime_groups.values()), "equipment_downtime_minutes": round(sum(row["minutes"] for row in downtime_rows), 2)}
    summary_text = f"Выдано {counts['issued']}, закрыто {counts['completed']}, просрочено на конец периода {counts['overdue']}, отклонено {counts['rejected']}. Подтверждённый простой оборудования — {counts['equipment_downtime_minutes']} мин. Найдено {len(anomalies)} сигналов для проверки мастером."
    return {"period": {"from_at": _iso(period.start), "to_at": _iso(period.end), "as_of": _iso(now), "semantics": "[from_at, to_at)", "timezone": "UTC; default display Asia/Qyzylorda", "shift_basis": "time_overlap_independent_of_issuing_shift" if period.shift_id else None}, "filters": {name: getattr(period, name) for name in ("section_id", "equipment_id", "employee_id", "brigade_id", "shift_id")}, "roster": {"shift_id": period.shift_id, "member_ids": sorted(visible_roster_ids), "basis": "ShiftMember", "includes_actual_activity_outside_roster": True}, "brigade_attribution": {"basis": "immutable_assignment_submission_event_snapshots", "unknown_execution_orders": sum(_execution_brigade(order) is None for order in orders), "estimates_used": False}, "scope": "own" if user.role == "worker" else "enterprise", "counts": counts, "order_ids": ids, "overdue_order_ids": [order.id for order in overdue], "workload": workload, "ratings": {"employees": rating_rows, "brigades": brigade_rows if user.role != "worker" else [], "formula": "40Q + 25T + 15F + 15C + 5A", "numeric_engine": "deterministic", "association_is_not_blame": True}, "downtime": downtime_rows, "materials": {"rows": material_rows, "totals": list(material_totals.values()), "basis": "latest_submission_snapshot_in_period"}, "anomalies": anomalies, "summary": {"text": summary_text, "engine": "rules", "model_status": "not_requested", "numeric_source": "deterministic_report"}, "synthetic": bool(orders) and all(order.synthetic for order in orders)}


def _anomalies(orders, history_orders, downtime_rows, material_rows, equipment, sections, faults, period, now) -> list[dict[str, Any]]:
    findings = []
    unplanned = [order for order in orders if order.work_type == "unplanned" and in_period(order.created_at, period.start, period.end) and as_utc(order.created_at) <= now]
    by_equipment = defaultdict(list)
    by_section = defaultdict(list)
    by_fault = defaultdict(list)
    by_hour = defaultdict(list)
    for order in unplanned:
        by_equipment[order.equipment_id].append(order)
        by_section[order.section_id].append(order)
        if order.fault_code_id:
            by_fault[(order.equipment_id, order.fault_code_id)].append(order)
        local_hour = as_utc(order.created_at).astimezone(timezone(timedelta(hours=5))).hour
        by_hour[local_hour // 6].append(order)
    for equipment_id, items in sorted(by_equipment.items(), key=lambda item: (-len(item[1]), item[0]))[:5]:
        if len(items) >= 3:
            findings.append({"kind": "frequent_failures", "equipment_id": equipment_id, "count": len(items), "text": f"{equipment[equipment_id].name}: {len(items)} внеплановых нарядов за выбранный период.", "recommendation": "Проверить причины повторных остановок и план обслуживания.", "evidence_order_ids": [order.id for order in items], "engine": "statistics", "association_is_not_blame": True})
    for section_id, items in sorted(by_section.items(), key=lambda item: (-len(item[1]), item[0]))[:5]:
        if len(items) >= 3:
            findings.append({"kind": "section_failures", "section_id": section_id, "count": len(items), "text": f"Участок {sections[section_id].name}: {len(items)} внеплановых нарядов.", "recommendation": "Сопоставить отказы с загрузкой оборудования и историей ремонтов.", "evidence_order_ids": [order.id for order in items], "engine": "statistics", "association_is_not_blame": True})
    for (equipment_id, fault_id), items in by_fault.items():
        if len(items) >= 3:
            findings.append({"kind": "repeated_fault", "equipment_id": equipment_id, "fault_code_id": fault_id, "count": len(items), "text": f"{equipment[equipment_id].name}: {len(items)} повторов шифра {faults[fault_id].code}.", "recommendation": "Проверить первопричину; повторение шифра само по себе не доказывает качество предыдущего ремонта.", "evidence_order_ids": [order.id for order in items], "engine": "statistics", "association_is_not_blame": True})
    for row in downtime_rows[:5]:
        if row["minutes"]:
            findings.append({"kind": "high_downtime", "equipment_id": row["equipment_id"], "minutes": row["minutes"], "text": f"{row['equipment_name']}: подтверждённый простой {row['minutes']} мин.", "recommendation": "Разобрать причины наиболее длительных остановок.", "evidence_order_ids": row["evidence_order_ids"], "engine": "statistics", "association_is_not_blame": True})
    for later in unplanned:
        previous = [order for order in history_orders if order.id != later.id and order.work_type == "planned" and order.equipment_id == later.equipment_id and order.closed_at and 0 <= (as_utc(later.created_at) - as_utc(order.closed_at)).total_seconds() <= 7 * 86400]
        if previous:
            nearest = max(previous, key=lambda order: as_utc(order.closed_at))
            findings.append({"kind": "failure_after_planned", "equipment_id": later.equipment_id, "text": "Внеплановый наряд создан в течение 7 дней после планового ремонта.", "recommendation": "Проверить связь событий и объём выполненной профилактики; автоматического обвинения исполнителя нет.", "evidence_order_ids": [nearest.id, later.id], "engine": "statistics", "association_is_not_blame": True})
    for row in material_rows:
        if row["norm_quantity"] is not None and row["quantity"] > row["norm_quantity"] * 1.5:
            findings.append({"kind": "material_over_norm", "material_id": row["material_id"], "quantity": row["quantity"], "norm_quantity": row["norm_quantity"], "text": f"Наряд №{row['number']}: расход {row['material_name']} {row['quantity']} {row['unit']} выше нормы {row['norm_quantity']}.", "recommendation": "Проверить списание и документированное обоснование дополнительного расхода.", "evidence_order_ids": [row["order_id"]], "engine": "rules", "association_is_not_blame": True})
    if len(unplanned) >= 10:
        for bucket, items in by_hour.items():
            if len(items) / len(unplanned) >= .6:
                findings.append({"kind": "time_association", "count": len(items), "sample_size": len(unplanned), "text": f"{len(items)} из {len(unplanned)} внеплановых нарядов приходится на {bucket * 6:02}:00–{(bucket + 1) * 6:02}:00 UTC+5.", "recommendation": "Сравнить часы работы и нагрузку; концентрация событий не доказывает причинную связь.", "evidence_order_ids": [order.id for order in items], "engine": "statistics", "association_is_not_blame": True})
        for dimension in ("shift_id", "executor_id", "brigade_id"):
            grouped = defaultdict(list)
            for order in unplanned:
                identifier = getattr(order, dimension)
                if identifier:
                    grouped[identifier].append(order)
            for identifier, items in grouped.items():
                if len(items) / len(unplanned) >= .6:
                    relevant_count = sum(getattr(order, dimension) == identifier for order in orders)
                    findings.append({"kind": "assignment_association", "dimension": dimension, "entity_id": identifier, "count": len(items), "assigned_order_count": relevant_count, "sample_size": len(unplanned), "text": f"В группе {dimension}={identifier} зарегистрировано {len(items)} из {len(unplanned)} внеплановых нарядов; всего назначенных нарядов в группе {relevant_count}.", "recommendation": "Сопоставить распределение работ, сложность и экспозицию оборудования. Связь с назначением не доказывает вину человека или бригады.", "evidence_order_ids": [order.id for order in items], "engine": "statistics", "association_is_not_blame": True})
    return findings


def report_sheets(report: dict[str, Any], kind: str = "shift") -> dict[str, list[dict[str, Any]]]:
    meta = [{**report["period"], **report["filters"], "scope": report["scope"], "synthetic": report["synthetic"], "brigade_attribution": report["brigade_attribution"], "shift_roster": report["roster"]}]
    result = {"Период и фильтры": meta}
    if kind == "shift":
        result.update({"Смена": [report["counts"]], "Загрузка": report["workload"], "Сводка": [report["summary"]]})
    if kind in {"shift", "ratings"}:
        result.update({"Исполнители": report["ratings"]["employees"], "Бригады": report["ratings"]["brigades"]})
    if kind in {"shift", "materials"}:
        result.update({"Материалы": report["materials"]["rows"], "Итого материалы": report["materials"]["totals"]})
    if kind in {"shift", "downtime"}:
        result["Простои"] = report["downtime"]
    if kind in {"shift", "anomalies"}:
        result["Сигналы для проверки"] = report["anomalies"]
    return result


def order_sheets(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    if report["audience"] == "worker":
        return {"Моя оценка": [{"number": report["number"], **report["quality"], **report["timing"]}]}
    return {"Наряд": [{key: value for key, value in report.items() if key not in {"history", "materials", "photos", "assessment"}}], "История": report["history"], "Материалы": report["materials"], "Фото": report["photos"], "Проверка": [report["assessment"]] if report["assessment"] else []}


def as_utc(value: datetime) -> datetime:
    """SQLite can return naive datetimes; persisted times are interpreted as UTC."""
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def in_period(value: datetime | None, start: datetime, end: datetime) -> bool:
    return value is not None and as_utc(start) <= as_utc(value) < as_utc(end)


def merged_intervals(
    intervals: Iterable[tuple[datetime, datetime | None]],
    start: datetime,
    end: datetime,
    now: datetime,
) -> list[tuple[datetime, datetime]]:
    """Clip real intervals to the period and merge them to avoid double counting."""
    period_start, period_end, current = as_utc(start), as_utc(end), as_utc(now)
    if period_end <= period_start:
        raise ValueError("Конец периода должен быть позже начала")
    clipped = []
    for interval_start, interval_end in intervals:
        left = max(as_utc(interval_start), period_start)
        right = min(as_utc(interval_end) if interval_end else current, period_end, current)
        if right > left:
            clipped.append((left, right))
    merged: list[tuple[datetime, datetime]] = []
    for left, right in sorted(clipped):
        if merged and left <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], right))
        else:
            merged.append((left, right))
    return merged


def interval_seconds(intervals: Iterable[tuple[datetime, datetime]]) -> float:
    return sum((as_utc(end) - as_utc(start)).total_seconds() for start, end in intervals)


def rating_components(components: Mapping[str, float | None]) -> dict[str, Any]:
    """Unknown evidence stays unknown; immature observations never earn free points."""
    available_weight = 0
    points = 0.0
    output = {}
    for name, weight in RATING_WEIGHTS.items():
        value = components.get(name)
        normalized = None if value is None else max(0.0, min(1.0, float(value)))
        earned = None if normalized is None else normalized * weight
        output[name] = {"weight": weight, "ratio": normalized, "points": None if earned is None else round(earned, 2)}
        if earned is not None:
            points += earned
            available_weight += weight
    return {
        "score": round(points / available_weight * 100, 2) if available_weight else None,
        "earned_points": round(points, 2),
        "available_weight": available_weight,
        "provisional": available_weight < 100,
        "components": output,
        "formula": "40Q + 25T + 15F + 15C + 5A",
    }


def safe_excel_cell(value: Any) -> Any:
    """Treat untrusted text as literal cells, including whitespace-prefixed formulas."""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, datetime):
        value = as_utc(value).isoformat()
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value).replace("\r\n", "\n").replace("\r", "\n")
        if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def export_workbook(sheets: Mapping[str, list[dict[str, Any]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name[:31])
        headers = list(dict.fromkeys(key for row in rows for key in row))
        if not headers:
            headers = ["Сообщение"]
            rows = [{"Сообщение": "Нет данных за выбранный период"}]
        sheet.append([safe_excel_cell(header) for header in headers])
        for row in rows:
            sheet.append([safe_excel_cell(row.get(header)) for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    if not workbook.worksheets:
        workbook.create_sheet("Отчёт").append(["Нет данных"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
